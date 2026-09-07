from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
import constants
import manuscript
from collections.abc import Callable



def create_document() -> Document:
    document = Document()

    section = document.sections[0]

    section.page_width = Inches(6)
    section.page_height = Inches(9)

    return document

def enable_mirrored_margins(
    document: Document,
) -> None:
    settings = document.settings.element

    if settings.find(qn("w:mirrorMargins")) is None:
        settings.append(
            OxmlElement("w:mirrorMargins")
        )

def apply_page_layout(
    document: Document,
    margin_profile: dict[str, float],
    gutter_width: float,
) -> None:
    section = document.sections[0]

    section.page_width = Inches(6)
    section.page_height = Inches(9)

    section.top_margin = Inches(
        margin_profile["top"]
    )

    section.bottom_margin = Inches(
        margin_profile["bottom"]
    )

    section.left_margin = Inches(
        margin_profile["outside"]
    )

    section.right_margin = Inches(
        margin_profile["outside"]
    )

    section.gutter = Inches(
        gutter_width
    )

    enable_mirrored_margins(
        document
    )

def apply_body_style(
    document: Document,
    font_name: str,
    font_size: float,
    line_spacing: float,
    first_line_indent: float,
) -> None:
    normal_style = document.styles["Normal"]

    normal_style.font.name = font_name
    normal_style.font.size = Pt(
        font_size
    )

    paragraph_format = (
        normal_style.paragraph_format
    )

    paragraph_format.line_spacing = (
        line_spacing
    )

    paragraph_format.first_line_indent = (
        Inches(first_line_indent)
    )

    paragraph_format.space_before = Pt(0)
    paragraph_format.space_after = Pt(0)

def add_text_page(
    document: Document,
    text: str,
    add_page_break: bool = True,
) -> None:
    for line in text.split("\n"):
        paragraph = document.add_paragraph()

        if line:
            paragraph.add_run(
                line
            )

    if add_page_break:
        document.add_page_break()

def add_title_page(
    document: Document,
    text: str,
    add_page_break: bool = True,
) -> None:
    lines = text.split("\n")

    for index, line in enumerate(lines):
        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph.paragraph_format.first_line_indent = None

        if index == 0:
            paragraph.paragraph_format.space_before = (
                Inches(2.0)
            )
            font_size = 20

        elif index == 1:
            paragraph.paragraph_format.space_after = (
                Inches(0.5)
            )
            font_size = 14

        else:
            font_size = 12

        run = paragraph.add_run(
            line
        )

        run.font.size = Pt(
            font_size
        )

    if add_page_break:
        document.add_page_break()

def add_copyright_page(
    document: Document,
    text: str,
    add_page_break: bool = True,
) -> None:
    lines = text.split("\n")

    for index, line in enumerate(lines):
        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.LEFT
        )

        paragraph.paragraph_format.first_line_indent = None
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)

        if index == 0:
            paragraph.paragraph_format.space_before = (
                Inches(1.5)
            )

        run = paragraph.add_run(
            line
        )

        run.font.size = Pt(
            8.0
        )

    if add_page_break:
        document.add_page_break()

def add_dedication_page(
    document: Document,
    text: str,
    add_page_break: bool = True,
) -> None:
    lines = text.split("\n")

    for index, line in enumerate(lines):
        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph.paragraph_format.first_line_indent = None
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)

        if index == 0:
            paragraph.paragraph_format.space_before = (
                Inches(2.5)
            )

        paragraph.add_run(
            line
        )

    if add_page_break:
        document.add_page_break()

def add_map_page(
    document: Document,
    image_path: str,
    add_page_break: bool = True,
) -> None:
    section = document.sections[0]

    usable_width = (
        section.page_width
        - section.left_margin
        - section.right_margin
        - section.gutter
    )

    paragraph = document.add_paragraph()

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    run = paragraph.add_run()

    run.add_picture(
        image_path,
        width=usable_width,
    )

    if add_page_break:
        document.add_page_break()

def add_trigger_warnings_page(
    document: Document,
    text: str,
    add_page_break: bool = True,
) -> None:
    lines = text.split("\n")

    for index, line in enumerate(lines):
        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.LEFT
        )

        paragraph.paragraph_format.first_line_indent = None
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)

        if index == 0:
            paragraph.paragraph_format.space_before = (
                Inches(1.0)
            )

        paragraph.add_run(
            line
        )

    if add_page_break:
        document.add_page_break()

# dispatcher
def add_page(
    document: Document,
    text: str = "",
    page_type: str = "body",
    image_path: str | None = None,
    add_page_break: bool = True,
) -> None:
    if page_type == "title_page":
        add_title_page(
            document,
            text,
            add_page_break,
        )

    elif page_type == "copyright":
        add_copyright_page(
            document,
            text,
            add_page_break,
        )

    elif page_type == "dedication":
        add_dedication_page(
            document,
            text,
            add_page_break,
        )

    elif page_type == "map":
        if image_path is None:
            raise ValueError(
                "Map page requires an image path."
            )

        add_map_page(
            document,
            image_path,
            add_page_break,
        )

    elif page_type == "trigger_warnings":
        add_trigger_warnings_page(
            document,
            text,
            add_page_break,
        )

    else:
        add_body_page(
            document,
            text,
            add_page_break,
        )

def export_document(
    pages: list[str],
    page_types: list[str],
    page_image_paths: list[str | None],
    save_path: str,
    margin_profile: dict[str, float],
    gutter_width: float,
    font_name: str,
    font_size: float,
    line_spacing: float,
    first_line_indent: float,
    include_page_numbers: bool,
    progress_callback: Callable[[int, int], None] | None = None,
) -> None:
    if not (
        len(pages)
        == len(page_types)
        == len(page_image_paths)
    ):
        raise ValueError(
            "Page data is out of sync."
        )

    document = create_document()

    apply_page_layout(
        document,
        margin_profile,
        gutter_width,
    )

    apply_body_style(
        document,
        font_name,
        font_size,
        line_spacing,
        first_line_indent,
    )

    body_start_index = next(
        (
            index
            for index, page_type
            in enumerate(page_types)
            if page_type == "body"
        ),
        None,
    )

    if (
        include_page_numbers
        and body_start_index == 0
    ):
        body_section = document.sections[0]

        restart_page_numbering(
            body_section,
            1,
        )

        add_page_number(
            body_section,
            font_name,
            font_size,
        )

    for index, (
        text,
        page_type,
        image_path,
    ) in enumerate(
        zip(
            pages,
            page_types,
            page_image_paths,
        )
    ):
        if (
            include_page_numbers
            and body_start_index is not None
            and body_start_index > 0
            and index == body_start_index
        ):
            body_section = document.add_section(
                WD_SECTION.NEW_PAGE
            )

            body_section.footer.is_linked_to_previous = (
                False
            )

            restart_page_numbering(
                body_section,
                1,
            )

            add_page_number(
                body_section,
                font_name,
                font_size,
            )

        is_last_page = (
            index == len(pages) - 1
        )

        next_page_starts_body = (
            include_page_numbers
            and body_start_index is not None
            and index == body_start_index - 1
        )

        add_page(
            document,
            text=text,
            page_type=page_type,
            image_path=image_path,
            add_page_break=(
                not is_last_page
                and not next_page_starts_body
            ),
        )

        document.save(
            save_path
        )

        if progress_callback is not None:
            progress_callback(
                index + 1,
                len(pages),
            )

def add_body_page(
    document: Document,
    text: str,
    add_page_break: bool = True,
) -> None:
    previous_was_heading = False

    for line in text.split("\n"):
        paragraph = document.add_paragraph()

        is_heading = manuscript.is_chapter_heading(
            line.strip()
        )

        if is_heading:
            print(
                "EXPORT CHAPTER HEADING:",
                repr(line),
            )

        if is_heading:
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            paragraph.paragraph_format.first_line_indent = (
                Inches(0)
            )

            paragraph.paragraph_format.space_after = (
                Inches(
                    constants.CHAPTER_HEADING_SPACE_AFTER_INCHES
                )
            )

            run = paragraph.add_run(
                line
            )

            run.font.size = Pt(
                constants.CHAPTER_HEADING_FONT_SIZE
            )

        else:
            if previous_was_heading:
                paragraph.paragraph_format.first_line_indent = (
                    Inches(0)
                )

            paragraph.add_run(
                line
            )

        previous_was_heading = is_heading

    if add_page_break:
        document.add_page_break()

def add_page_number(
    section,
    font_name: str,
    font_size: float,
) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = paragraph.add_run()
    run.font.name = font_name

    run.font.size = Pt(
        font_size
    )   

    field_begin = OxmlElement("w:fldChar")
    field_begin.set(
        qn("w:fldCharType"),
        "begin",
    )

    field_code = OxmlElement("w:instrText")
    field_code.set(
        qn("xml:space"),
        "preserve",
    )
    field_code.text = " PAGE "

    field_end = OxmlElement("w:fldChar")
    field_end.set(
        qn("w:fldCharType"),
        "end",
    )

    run._r.append(field_begin)
    run._r.append(field_code)
    run._r.append(field_end)


def restart_page_numbering(
    section,
    start: int = 1,
) -> None:
    section_properties = section._sectPr

    page_number_type = (
        section_properties.find(
            qn("w:pgNumType")
        )
    )

    if page_number_type is None:
        page_number_type = OxmlElement(
            "w:pgNumType"
        )
        section_properties.append(
            page_number_type
        )

    page_number_type.set(
        qn("w:start"),
        str(start),
    )